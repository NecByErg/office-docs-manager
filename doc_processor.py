from docx import Document
from unicode_to_preeti import to_preeti

def process_document(template_path, output_path, data):
    doc = Document(template_path)
    
    # Text Replace Logic
    for p in doc.paragraphs:
        if "{{ADDRESS_BLOCK}}" in p.text or "Address Block" in p.text:
            p.text = ""  # Paragraph सफा गर्ने
            
            # Form बाट आएको Multi-line address
            lines = data.get("address_block", "").split("\n")
            
            for line in lines:
                # Unicode लाई Preeti ASCII मा रूपान्तरण गर्ने
                preeti_text = to_preeti(line.strip())
                
                # Dynamic Line append गर्ने
                run = p.add_run(preeti_text + "\n")
                run.font.name = "Preeti"  # Preeti Font Set गर्ने
                
    doc.save(output_path)
